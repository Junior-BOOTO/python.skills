"""Exports Word/PDF du retour pédagogique validé (sans nom d'apprenant)."""
from __future__ import annotations
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape
from .engine import Analysis, review_suggestions


def _entries(analysis: Analysis, selected: set[int]) -> list:
    return [s for i, s in enumerate(analysis.suggestions) if i in selected]


def to_text(analysis: Analysis, selected: set[int], note: str = "") -> str:
    corrected = review_suggestions(analysis.original, analysis.suggestions, selected)
    lines = ["BOOTO Correcteur FLE — Retour pédagogique", "Prototype A1 — propositions validées par l'enseignant", "", "TEXTE INITIAL", analysis.original, "", "TEXTE RÉVISÉ", corrected, "", "PISTES RETENUES"]
    for s in _entries(analysis, selected):
        lines.append(f"• {s.original} → {s.proposal} ({s.category}) : {s.explanation}")
    if not selected:
        lines.append("Aucune proposition retenue.")
    lines += ["", "OBSERVATION ENSEIGNANT", note.strip() or "Aucune observation renseignée.", "", "Attention : aucune note ni certification CECRL n'est attribuée par le moteur."]
    return "\n".join(lines)


def to_docx(analysis: Analysis, selected: set[int], note: str = "") -> bytes:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    d = Document()
    d.core_properties.title = "BOOTO Correcteur FLE — Retour pédagogique"
    d.core_properties.author = "Junior BOOTO Waba"
    sec=d.sections[0]
    sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2.1);sec.right_margin=Cm(2.1)
    normal=d.styles['Normal'];normal.font.name='Aptos';normal.font.size=Pt(10);normal.paragraph_format.space_after=Pt(6)
    title=d.add_heading('BOOTO Correcteur FLE',0)
    title.runs[0].font.color.rgb=RGBColor(21,72,120)
    d.add_paragraph('Retour pédagogique · A1 · Propositions examinées par l’enseignant')
    d.add_heading('Texte initial',level=1);d.add_paragraph(analysis.original)
    d.add_heading('Version révisée',level=1);d.add_paragraph(review_suggestions(analysis.original,analysis.suggestions,selected))
    d.add_heading('Pistes de révision retenues',level=1)
    entries = _entries(analysis,selected)
    for s in entries:
        d.add_paragraph(f'{s.original} → {s.proposal} — {s.explanation}',style='List Bullet')
    if not entries:d.add_paragraph('Aucune proposition retenue.')
    d.add_heading('Observations de l’enseignant',level=1);d.add_paragraph(note.strip() or 'Aucune observation renseignée.')
    p=d.add_paragraph('Ce document ne constitue ni une correction exhaustive, ni une note, ni une certification CECRL.')
    p.runs[0].italic=True
    f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;f.text='BOOTO AI & DATA INNOVATION LAB · © 2026 Junior BOOTO Waba'
    stream=BytesIO();d.save(stream);return stream.getvalue()


def to_pdf(analysis: Analysis, selected: set[int], note: str = "") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    font='Helvetica'
    for path in (Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'),Path('/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf')):
        if path.exists():
            font='CorrecteurFLEFont'
            if font not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont(font,str(path)))
            break
    h=ParagraphStyle('h',fontName=font,fontSize=17,leading=23,textColor=colors.HexColor('#154878'),spaceAfter=9)
    section=ParagraphStyle('section',fontName=font,fontSize=11,leading=16,textColor=colors.HexColor('#078f91'),spaceBefore=10,spaceAfter=6)
    body=ParagraphStyle('body',fontName=font,fontSize=9,leading=15,spaceAfter=6)
    small=ParagraphStyle('small',fontName=font,fontSize=7.8,leading=12,textColor=colors.HexColor('#516980'),spaceAfter=8)
    def par(text:str,style):
        return Paragraph(escape(text).replace('\n','<br/>'),style)
    story=[par('BOOTO Correcteur FLE',h),par('Retour pédagogique · A1 · Propositions examinées par l’enseignant',small),par('Texte initial',section),par(analysis.original,body),par('Version révisée',section),par(review_suggestions(analysis.original,analysis.suggestions,selected),body),par('Pistes de révision retenues',section)]
    entries=_entries(analysis,selected)
    if entries:
        for s in entries:story.append(KeepTogether([par(f'{s.original} → {s.proposal}',body),par(s.explanation,small)]))
    else:story.append(par('Aucune proposition retenue.',body))
    story.extend([par('Observations de l’enseignant',section),par(note.strip() or 'Aucune observation renseignée.',body),Spacer(1,6*mm),par('Ce document ne constitue ni une correction exhaustive, ni une note, ni une certification CECRL.',small)])
    target=BytesIO()
    doc=SimpleDocTemplate(target,pagesize=A4,leftMargin=20*mm,rightMargin=20*mm,topMargin=18*mm,bottomMargin=18*mm,title='BOOTO Correcteur FLE — Retour pédagogique',author='Junior BOOTO Waba')
    def footer(canv,doc):
        canv.saveState();canv.setFont(font,7);canv.setFillColor(colors.HexColor('#5a7183'));canv.drawCentredString(A4[0]/2,11*mm,f'BOOTO Correcteur FLE · page {doc.page}');canv.restoreState()
    doc.build(story,onFirstPage=footer,onLaterPages=footer)
    return target.getvalue()

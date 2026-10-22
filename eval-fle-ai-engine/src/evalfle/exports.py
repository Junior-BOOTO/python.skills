"""Exports imprimables en Word et PDF, destinés à l'enseignant."""
from __future__ import annotations
from io import BytesIO
from xml.sax.saxutils import escape

from .engine import Question, overview, safe_label


def _title(grade: int):
    return f'Évaluation de français — Grade {grade} — A1 (supports illustratifs)'


def build_docx(questions: list[Question], teacher: bool = False, author: str = 'Junior BOOTO Waba') -> bytes:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    if not questions:
        raise ValueError('Questions manquantes')
    grade = questions[0].grade
    d=Document()
    sec=d.sections[0];sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
    styles=d.styles
    normal=styles['Normal'];normal.font.name='Aptos';normal.font.size=Pt(10);normal.font.color.rgb=RGBColor(29,50,80)
    normal.paragraph_format.space_after=Pt(5)
    for name,sz in [('Title',18),('Heading 1',13),('Heading 2',11)]:
        st=styles[name];st.font.name='Aptos';st.font.size=Pt(sz);st.font.color.rgb=RGBColor(21,91,142)
    p=d.add_paragraph(style='Title');p.add_run(_title(grade))
    if not teacher:
        p=d.add_paragraph('Nom : __________________________________________     Date : __________________')
        p.add_run('\nClasse : ______________________     Durée indicative : 45 minutes')
    label = ('CORRIGÉ ENSEIGNANT — RÉPONSES MARQUÉES' if teacher else 'DOCUMENT ÉLÈVE — QUESTIONNAIRE')
    p = d.add_paragraph()
    run = p.add_run(f'{label}  |  {len(questions)} QCM')
    run.bold = True
    run.font.color.rgb = RGBColor(15,113,92) if teacher else RGBColor(21,91,142)
    if teacher:
        d.add_paragraph('Les bonnes réponses sont cochées et indiquées en vert. Cette version est réservée à l’enseignant.')
        d.add_heading('Grille de correction rapide',level=2)
        quick = d.add_table(rows=0, cols=3)
        quick.style = 'Table Grid'
        chunk = (len(questions)+2)//3
        for i in range(chunk):
            cells = quick.add_row().cells
            for column in range(3):
                idx = i + column*chunk
                cells[column].text = f'{idx+1:02d}. {questions[idx].answer}' if idx < len(questions) else ''
                if idx < len(questions):
                    for r in cells[column].paragraphs[0].runs:
                        r.bold = True
                        r.font.color.rgb = RGBColor(15,113,92)
        d.add_heading('Questionnaire corrigé',level=2)
    else:
        d.add_paragraph('Consigne : coche une seule réponse (A, B, C ou D) pour chaque question. Aucun passage de lecture, enregistrement ou tâche orale.')
    for k,q in enumerate(questions,1):
        p=d.add_paragraph();p.paragraph_format.keep_with_next=True
        run=p.add_run(f'{k}. {q.question}');run.bold=True
        for letter,choice in zip('ABCD',q.choices):
            p=d.add_paragraph(style='Normal');p.paragraph_format.left_indent=Inches(.23)
            r = p.add_run(f'{"☑" if teacher and letter == q.answer else "☐"} {letter}. {choice}')
            if teacher and letter == q.answer:
                r.bold = True
                r.font.color.rgb = RGBColor(15,113,92)
    d.add_page_break()
    d.add_heading('Grille de correction cochée' if teacher else 'Grille de réponses',level=1)
    t=d.add_table(rows=1,cols=5)
    t.style='Table Grid'
    for c,txt in zip(t.rows[0].cells,['N°','A','B','C','D']):c.text=txt
    for k in range(1,len(questions)+1):
        row=t.add_row().cells;row[0].text=str(k)
        for i in range(1,5):
            marked = teacher and "ABCD"[i-1] == questions[k-1].answer
            row[i].text = 'X' if marked else '☐'
            if marked:
                for r in row[i].paragraphs[0].runs:
                    r.bold = True
                    r.font.color.rgb = RGBColor(15,113,92)
    if teacher:
        d.add_page_break()
        d.add_heading('Corrigé réservé à l’enseignant',level=1)
        for k,q in enumerate(questions,1):
            d.add_paragraph(f'{k:02d} — {q.answer}  |  {q.category}  |  {q.choices["ABCD".index(q.answer)]}')
        d.add_heading('Critères observables « Je peux… »',level=2)
        for row in overview(questions):
            d.add_paragraph(f"• {row['category']} ({row['count']} items) : {row['criterion']}")
        d.add_paragraph('Barème : 1 point par réponse correcte, 0 pour réponse incorrecte ou absente. Résultats par domaine à interpréter en contexte ; ce questionnaire ne certifie aucun niveau du CECRL.')
        d.add_paragraph('Limite : les QCM ne mesurent pas directement la compréhension orale, la production orale ou la production écrite.')
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.text=f'BOOTO Éval’FLE AI Engine · Prototype V1 · {safe_label(author)}'
    b=BytesIO();d.save(b);return b.getvalue()


def build_pdf(questions: list[Question], teacher: bool = False, author: str = 'Junior BOOTO Waba') -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from pathlib import Path
    if not questions:
        raise ValueError('Questions manquantes')
    fonts=[Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'),Path('/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf')]
    chosen=next((p for p in fonts if p.exists()),None)
    font='Helvetica'
    if chosen:
        font='EvalFLEUnicode'
        if font not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont(font,str(chosen)))
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='EvalTitle',fontName=font,fontSize=14,leading=18,textColor=colors.HexColor('#144e82'),spaceAfter=12))
    styles.add(ParagraphStyle(name='EvalInfo',fontName=font,fontSize=8.6,leading=13,spaceAfter=7))
    styles.add(ParagraphStyle(name='EvalItem',fontName=font,fontSize=9.0,leading=13,spaceAfter=4))
    styles.add(ParagraphStyle(name='EvalOption',fontName=font,fontSize=8.7,leading=12,leftIndent=12,spaceAfter=2))
    styles.add(ParagraphStyle(name='EvalH2',fontName=font,fontSize=12,leading=15,textColor=colors.HexColor('#147c91'),spaceBefore=7,spaceAfter=8))
    stream=BytesIO();pdf=SimpleDocTemplate(stream,pagesize=A4,rightMargin=18*mm,leftMargin=18*mm,topMargin=17*mm,bottomMargin=17*mm)
    story=[Paragraph(escape(_title(questions[0].grade)),styles['EvalTitle'])]
    if teacher:
        story.extend([
            Paragraph('<font color="#0f715c"><b>CORRIGÉ ENSEIGNANT — RÉPONSES MARQUÉES</b></font>',styles['EvalH2']),
            Paragraph('Les bonnes réponses sont cochées et indiquées en vert. Réservé à l’enseignant.',styles['EvalInfo']),
            Paragraph('Grille de correction rapide',styles['EvalH2']),
        ])
        chunk_size=(len(questions)+2)//3
        quick_rows=[]
        for i in range(chunk_size):
            quick_rows.append([
                f'{idx+1:02d}. {questions[idx].answer}' if (idx:=i+column*chunk_size)<len(questions) else ''
                for column in range(3)
            ])
        quick=Table(quick_rows,colWidths=[42*mm]*3,rowHeights=[6*mm]*len(quick_rows))
        quick.setStyle(TableStyle([
            ('FONTNAME',(0,0),(-1,-1),font),('FONTSIZE',(0,0),(-1,-1),9),
            ('TEXTCOLOR',(0,0),(-1,-1),colors.HexColor('#0f715c')),
            ('GRID',(0,0),(-1,-1),.4,colors.HexColor('#b1d9d1')),
            ('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#edf8f4')),
            ('ALIGN',(0,0),(-1,-1),'CENTER'),('VALIGN',(0,0),(-1,-1),'MIDDLE')]))
        story.extend([quick,Spacer(1,4*mm),Paragraph('Questionnaire corrigé',styles['EvalH2'])])
    else:
        story.extend([
            Paragraph('Nom : ____________________________________      Date : ____________________',styles['EvalInfo']),
            Paragraph(f'Document élève — questionnaire | {len(questions)} QCM | Durée indicative : 45 minutes',styles['EvalInfo']),
            Paragraph('Consigne : coche une seule réponse (A, B, C ou D) par question. Chaque réponse correcte vaut 1 point.',styles['EvalInfo']),Spacer(1,4*mm)])
    for k,q in enumerate(questions,1):
        chunk=[Paragraph(f'<b>{k}. {escape(q.question)}</b>',styles['EvalItem'])]
        for letter,choice in zip('ABCD',q.choices):
            opt = f'(X) {letter}. {escape(choice)}' if teacher and letter==q.answer else f'(  ) {letter}. {escape(choice)}'
            chunk.append(Paragraph(f'<font color="#0f715c"><b>{opt}</b></font>' if teacher and letter==q.answer else opt,styles['EvalOption']))
        chunk.append(Spacer(1,3*mm))
        story.append(KeepTogether(chunk))
    story.extend([PageBreak(),Paragraph('Grille de correction cochée' if teacher else 'Grille de réponses',styles['EvalH2'])])
    rows=[['N°','A','B','C','D']]+[
        [str(k)]+[('(X)' if teacher and letter==questions[k-1].answer else '(  )') for letter in 'ABCD']
        for k in range(1,len(questions)+1)]
    t=Table(rows,colWidths=[22*mm]+[24*mm]*4,rowHeights=[8*mm]+[6.5*mm]*len(questions),repeatRows=1)
    fmt=[('FONTNAME',(0,0),(-1,-1),font),('FONTSIZE',(0,0),(-1,-1),8),('GRID',(0,0),(-1,-1),.4,colors.HexColor('#c0d4e6')),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e7f4fb')),('ALIGN',(0,0),(-1,-1),'CENTER'),('VALIGN',(0,0),(-1,-1),'MIDDLE')]
    if teacher:
        for k,q in enumerate(questions,1):
            cell=('ABCD'.index(q.answer)+1,k)
            fmt.extend([('BACKGROUND',cell,cell,colors.HexColor('#ddf4ea')),('TEXTCOLOR',cell,cell,colors.HexColor('#0f715c'))])
    t.setStyle(TableStyle(fmt))
    story.append(t)
    if teacher:
        story.extend([PageBreak(),Paragraph('Corrigé réservé à l’enseignant',styles['EvalH2'])])
        for k,q in enumerate(questions,1):
            story.append(Paragraph(f'{k:02d}. <b>{q.answer}</b> — {escape(q.category)} : {escape(q.choices["ABCD".index(q.answer)])}',styles['EvalItem']))
        story.append(Paragraph('Critères observables « Je peux… »',styles['EvalH2']))
        for row in overview(questions):
            story.append(Paragraph(f"<b>{escape(row['category'])} ({row['count']} items) :</b> {escape(row['criterion'])}",styles['EvalInfo']))
        story.append(Paragraph('Barème : 1 point par bonne réponse, 0 sinon. Ces résultats ne certifient pas le niveau CECRL ; les compétences orales et la production écrite ne sont pas évaluées directement.',styles['EvalInfo']))
    def footer(canvas,doc):
        canvas.saveState();canvas.setFont(font,7);canvas.setFillColor(colors.HexColor('#54708d'))
        canvas.drawCentredString(A4[0]/2,10*mm,f"BOOTO Éval’FLE AI Engine · {safe_label(author)} · page {doc.page}")
        canvas.restoreState()
    pdf.build(story,onFirstPage=footer,onLaterPages=footer)
    return stream.getvalue()
